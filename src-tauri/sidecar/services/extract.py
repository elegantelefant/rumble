# ABOUTME: Text extraction from uploaded documents for the review path.
# ABOUTME: Handles PDF and DOCX; fails honestly when there is no text layer.

import io
import logging
import zipfile

from docx import Document
from docx.table import _Cell
from pypdf import PasswordType, PdfReader

logger = logging.getLogger(__name__)

UNREADABLE_PDF = "This PDF could not be read. It may be damaged or not a real PDF."
UNREADABLE_DOCX = (
    "This Word document could not be read. It may be damaged, or saved in the "
    "older .doc format rather than .docx."
)
PASSWORD_PROTECTED_PDF = (
    "This PDF is password protected. Remove the password and upload it again."
)
DOCX_TOO_LARGE = "This Word document is too large or complex to process."

# A .docx is a zip, and python-docx inflates every part into memory and builds
# an XML tree from each XML part: a 232 KB upload of repetitive XML took 1.6 GB
# and 69 s. Two bounds, both checked before python-docx runs:
# - bytes: zipfile stops each member at its declared size, so the declared sizes
#   bound what python-docx can inflate. A document of images totalling more is refused.
# - nodes: the tree costs ~300-400 bytes per element or attribute whatever its
#   size; every element needs a "<" and every attribute a "=", so counting both
#   bounds it. Word-authored documents run to ~70 per paragraph, so this admits
#   ~20,000 paragraphs, several hundred pages.
MAX_DOCX_UNCOMPRESSED_MB = 32
MAX_DOCX_UNCOMPRESSED_BYTES = MAX_DOCX_UNCOMPRESSED_MB * 1024 * 1024
MAX_DOCX_XML_NODES = 1_500_000
_INFLATE_CHUNK_BYTES = 1024 * 1024


class ExtractionError(Exception):
    """Raised when a document cannot be turned into usable text."""


def _pdf_pages(data: bytes) -> list[str]:
    reader = PdfReader(io.BytesIO(data))

    # Print/copy-restricted PDFs carry only an owner password and open with an
    # empty user password; only a real user password blocks reading.
    if reader.is_encrypted and reader.decrypt("") == PasswordType.NOT_DECRYPTED:
        raise ExtractionError(PASSWORD_PROTECTED_PDF)

    return [page.extract_text() or "" for page in reader.pages]


def extract_pdf(data: bytes) -> str:
    # pypdf parses lazily, so a damaged file can fail at open, at reader.pages or
    # in extract_text(), with anything from PdfReadError to KeyError.
    try:
        pages = _pdf_pages(data)
    except ExtractionError:
        raise
    except Exception as exc:
        logger.warning("PDF parse failed: %s: %s", type(exc).__name__, exc)
        raise ExtractionError(UNREADABLE_PDF) from exc

    text = "\n\n".join(p for p in pages if p.strip())

    if not text.strip():
        # A scanned document is images with no text layer. pypdf returns empty
        # strings rather than failing, so this is the only signal we get.
        raise ExtractionError(
            "No text found in this PDF. It may be a scan or image-only document."
        )

    return text


def _docx_over_limit(archive: zipfile.ZipFile) -> str | None:
    """Which bound the archive exceeds, if any; inflates in chunks only once the byte bound holds."""
    members = archive.infolist()
    uncompressed = sum(member.file_size for member in members)
    if uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
        return f"{uncompressed} bytes uncompressed"
    nodes = 0
    for member in members:
        with archive.open(member) as part:
            while chunk := part.read(_INFLATE_CHUNK_BYTES):
                nodes += chunk.count(b"<") + chunk.count(b"=")
    return f"{nodes} XML elements and attributes" if nodes > MAX_DOCX_XML_NODES else None


def _docx_parts(data: bytes) -> list[str]:
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            # The row's own <w:tc> cells, not row.cells: that repeats a cell once
            # per grid column it spans and walks up the table for each vertically
            # merged one, so a crafted gridSpan ran past the host timeout (and
            # merged text came out repeated).
            cells = [text for tc in row._tr.tc_lst if (text := _Cell(tc, table).text.strip())]
            if cells:
                parts.append(" | ".join(cells))
    return parts


def extract_docx(data: bytes) -> str:
    # python-docx builds its proxies lazily, so as for PDFs the whole
    # extraction sits inside the guard, not just the open.
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            over_limit = _docx_over_limit(archive)
        if over_limit:
            logger.warning("DOCX refused: %s", over_limit)
            raise ExtractionError(DOCX_TOO_LARGE)
        parts = _docx_parts(data)
    except ExtractionError:
        raise
    except Exception as exc:
        logger.warning("DOCX parse failed: %s: %s", type(exc).__name__, exc)
        raise ExtractionError(UNREADABLE_DOCX) from exc

    text = "\n\n".join(parts)

    if not text.strip():
        raise ExtractionError("No text found in this document.")

    return text


def extract(filename: str, data: bytes) -> str:
    """Dispatch on file extension. Plain text passes through unchanged."""
    lower = filename.lower()

    if lower.endswith(".pdf"):
        return extract_pdf(data)
    if lower.endswith(".docx"):
        return extract_docx(data)
    if lower.endswith((".txt", ".md")):
        try:
            return data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ExtractionError("File is not valid UTF-8 text.") from exc

    raise ExtractionError(f"Unsupported file type: {filename}")
